"""
src/resume_customizer.py - JD-aware resume customizer for Siddhu and future clients.

Existing main.py compatibility:
    customize_resume(post_text, lead_index) -> generated PDF path

Only placeholder content is changed. Experience, projects, education, dates,
layout, and existing resume formatting remain untouched.
"""

import json
import re
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from docx2pdf import convert

from config import BASE_DIR, CANDIDATE_DATA_FILE


TEMPLATES_DIR = BASE_DIR / "templates"
OUTPUT_RESUME_DIR = BASE_DIR / "outputs" / "resumes"

MISSING_VALUES = {
    "", "none", "nan", "n/a", "na", "null", "not available",
    "not avail", "not mentioned", "unknown",
}


# -----------------------------------------------------------------------------
# Generic helpers
# -----------------------------------------------------------------------------

def safe_filename(value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"[^A-Za-z0-9]+", "_", text)
    return re.sub(r"_+", "_", text).strip("_") or "Candidate"


def clean_text(value: Any) -> str:
    """Produce a single paragraph with exactly one space between words."""
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ").replace("\t", " ")
    return re.sub(r"\s+", " ", text).strip()


def ordered_unique(items: List[str]) -> List[str]:
    output: List[str] = []
    seen = set()
    for item in items:
        item = clean_text(item)
        if not item:
            continue
        key = item.lower()
        if key not in seen:
            seen.add(key)
            output.append(item)
    return output


def load_candidate_data() -> Dict[str, Any]:
    fallback = {
        "full_name": "Siddhu Kolamala",
        "name": "Siddhu Kolamala",
        "title": "Digital Marketing Specialist",
        "email": "supaysid@gmail.com",
        "phone": "+1 973-687-9494",
        "linkedin": "http://linkedin.com/in/kolamala-siddhu",
        "availability": "Immediate",
        "location": "New York, NY",
        "current_location": "New York, NY",
        "work_authorization": "",
        "total_experience": "",
        "open_to_relocate": "Yes",
    }

    path = Path(CANDIDATE_DATA_FILE)
    if not path.exists():
        return fallback

    with open(path, "r", encoding="utf-8") as file:
        supplied = json.load(file) or {}

    fallback.update(supplied)
    return fallback


def get_candidate_value(candidate: Dict[str, Any], *keys: str, default: str = "") -> str:
    for key in keys:
        value = clean_text(candidate.get(key, ""))
        if value and value.lower() not in MISSING_VALUES:
            return value
    return default


def jd_has(post_text: Any, terms: List[str]) -> bool:
    text = clean_text(post_text).lower()
    return any(term.lower() in text for term in terms)


# -----------------------------------------------------------------------------
# Template detection
# -----------------------------------------------------------------------------

def detect_resume_template(candidate: Dict[str, Any]) -> Path:
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)

    explicit = get_candidate_value(candidate, "resume_template", default="")
    if explicit:
        explicit_path = Path(explicit)
        if not explicit_path.is_absolute():
            explicit_path = BASE_DIR / explicit_path
        if explicit_path.exists():
            return explicit_path

    name = get_candidate_value(candidate, "full_name", "name", default="Candidate")
    safe_name = safe_filename(name)

    preferred = [
        TEMPLATES_DIR / f"{safe_name}_Resume.docx",
        TEMPLATES_DIR / f"{safe_name}.docx",
        TEMPLATES_DIR / "Resume_Template.docx",
        TEMPLATES_DIR / "resume_template.docx",
    ]

    for path in preferred:
        if path.exists():
            return path

    matching = list(TEMPLATES_DIR.glob(f"*{safe_name}*.docx"))
    if matching:
        return matching[0]

    all_templates = list(TEMPLATES_DIR.glob("*.docx"))
    if len(all_templates) == 1:
        return all_templates[0]

    raise FileNotFoundError(
        "Resume template not found. Add either "
        f"templates/{safe_name}_Resume.docx or templates/{safe_name}.docx"
    )


# -----------------------------------------------------------------------------
# Role and skill knowledge base
# -----------------------------------------------------------------------------

ROLE_RULES = [
    {
        "title": "SEO SPECIALIST",
        "category": "seo",
        "triggers": [
            "seo specialist", "seo executive", "seo analyst", "search engine optimization",
            "keyword research", "on-page seo", "off-page seo", "technical seo",
            "google search console", "backlink", "organic traffic",
        ],
    },
    {
        "title": "PPC / GOOGLE ADS SPECIALIST",
        "category": "ppc",
        "triggers": [
            "ppc", "google ads", "adwords", "paid search", "sem", "search ads",
            "display ads", "shopping ads", "cpc", "ctr", "conversion tracking",
        ],
    },
    {
        "title": "EMAIL MARKETING SPECIALIST",
        "category": "email_marketing",
        "triggers": [
            "email marketing", "mailchimp", "hubspot", "klaviyo", "drip campaign",
            "newsletter", "email automation", "marketing automation",
        ],
    },
    {
        "title": "SOCIAL MEDIA MARKETING SPECIALIST",
        "category": "social_media",
        "triggers": [
            "social media marketing", "social media specialist", "meta ads", "facebook ads",
            "instagram marketing", "linkedin ads", "community management",
        ],
    },
    {
        "title": "DIGITAL MARKETING SPECIALIST",
        "category": "digital_marketing",
        "triggers": [
            "digital marketing", "digital marketer", "performance marketing", "growth marketing",
            "content marketing", "lead generation", "demand generation", "campaign management",
            "google analytics", "marketing specialist", "marketing executive",
        ],
    },
]


BASE_SKILLS = {
    "digital_marketing": [
        "Digital Marketing Strategy", "SEO", "Search Engine Marketing", "Keyword Research",
        "On-Page SEO", "Off-Page SEO", "Technical SEO", "Google Search Console",
        "Google Analytics", "GA4", "Google Ads", "PPC Campaign Management", "Paid Search",
        "Performance Marketing", "Campaign Management", "Campaign Optimization",
        "Landing Page Optimization", "Conversion Rate Optimization", "Lead Generation",
        "Demand Generation", "Email Marketing", "Mailchimp", "HubSpot",
        "Marketing Automation", "Social Media Marketing", "Content Strategy",
        "Content Marketing", "Competitor Analysis", "A/B Testing", "Campaign Reporting",
    ],
    "seo": [
        "SEO", "Search Engine Optimization", "Keyword Research", "On-Page SEO",
        "Off-Page SEO", "Technical SEO", "SEO Audits", "Google Search Console",
        "Google Analytics", "GA4", "Content Optimization", "Meta Tag Optimization",
        "Internal Linking", "Backlink Analysis", "Organic Traffic Growth",
        "SERP Analysis", "Competitor Analysis", "Landing Page Optimization",
        "Conversion Rate Optimization", "WordPress", "Content Strategy",
        "Campaign Reporting", "A/B Testing", "Lead Generation",
    ],
    "ppc": [
        "Google Ads", "PPC Campaign Management", "Paid Search", "SEM",
        "Search Campaigns", "Display Campaigns", "Campaign Optimization",
        "Keyword Research", "Negative Keyword Management", "Ad Copy Optimization",
        "Conversion Tracking", "Landing Page Optimization", "CPC Optimization",
        "CTR Improvement", "Quality Score Optimization", "Budget Management",
        "Google Analytics", "GA4", "Performance Reporting", "A/B Testing",
        "Lead Generation", "Conversion Rate Optimization",
    ],
    "email_marketing": [
        "Email Marketing", "Mailchimp", "HubSpot", "Marketing Automation",
        "Email Campaign Management", "Campaign Segmentation", "Lead Nurturing",
        "Newsletter Campaigns", "Drip Campaigns", "Email Automation",
        "A/B Testing", "Open Rate Optimization", "Click-Through Rate Optimization",
        "Campaign Reporting", "Content Strategy", "CRM Coordination",
        "Landing Page Optimization", "Lead Generation", "Conversion Optimization",
    ],
    "social_media": [
        "Social Media Marketing", "Meta Ads", "Facebook Ads", "Instagram Marketing",
        "LinkedIn Ads", "Content Planning", "Content Calendar Management",
        "Community Engagement", "Audience Targeting", "Campaign Management",
        "Paid Social Campaigns", "Lead Generation", "Performance Reporting",
        "A/B Testing", "Conversion Tracking", "Brand Awareness",
    ],
}

JD_KEYWORD_SKILLS = {
    "seo": "SEO",
    "search engine optimization": "Search Engine Optimization",
    "keyword research": "Keyword Research",
    "on-page": "On-Page SEO",
    "on page": "On-Page SEO",
    "off-page": "Off-Page SEO",
    "off page": "Off-Page SEO",
    "technical seo": "Technical SEO",
    "seo audit": "SEO Audits",
    "google search console": "Google Search Console",
    "gsc": "Google Search Console",
    "google analytics": "Google Analytics",
    "ga4": "GA4",
    "serp": "SERP Analysis",
    "backlink": "Backlink Analysis",
    "organic traffic": "Organic Traffic Growth",
    "wordpress": "WordPress",
    "ppc": "PPC Campaign Management",
    "google ads": "Google Ads",
    "adwords": "Google Ads",
    "paid search": "Paid Search",
    "sem": "Search Engine Marketing",
    "meta ads": "Meta Ads",
    "facebook ads": "Facebook Ads",
    "instagram": "Instagram Marketing",
    "linkedin ads": "LinkedIn Ads",
    "performance marketing": "Performance Marketing",
    "growth marketing": "Growth Marketing",
    "email marketing": "Email Marketing",
    "mailchimp": "Mailchimp",
    "hubspot": "HubSpot",
    "klaviyo": "Klaviyo",
    "marketing automation": "Marketing Automation",
    "lead nurturing": "Lead Nurturing",
    "newsletter": "Newsletter Campaigns",
    "drip campaign": "Drip Campaigns",
    "social media": "Social Media Marketing",
    "content marketing": "Content Marketing",
    "content strategy": "Content Strategy",
    "lead generation": "Lead Generation",
    "demand generation": "Demand Generation",
    "landing page": "Landing Page Optimization",
    "conversion rate": "Conversion Rate Optimization",
    "cro": "Conversion Rate Optimization",
    "campaign management": "Campaign Management",
    "campaign optimization": "Campaign Optimization",
    "competitor analysis": "Competitor Analysis",
    "a/b testing": "A/B Testing",
    "reporting": "Campaign Reporting",
    "analytics": "Marketing Analytics",
    "crm": "CRM Coordination",
    "canva": "Canva",
    "copywriting": "Marketing Copywriting",
    "budget": "Campaign Budget Management",
    "conversion tracking": "Conversion Tracking",
    "content calendar": "Content Calendar Management",
}

SOFT_SKILLS = [
    "Communication", "Cross-functional Collaboration", "Analytical Problem Solving",
    "Stakeholder Communication", "Time Management",
]


# -----------------------------------------------------------------------------
# Role detection and JD analysis
# -----------------------------------------------------------------------------

def detect_role(post_text: Any, candidate: Dict[str, Any]) -> Dict[str, Any]:
    """
    Candidate title is intentionally given strong priority. This prevents an
    unrelated SOC/Java/other post from turning Siddhu's resume into another role.
    """
    candidate_title = get_candidate_value(candidate, "title", default="Digital Marketing Specialist")
    candidate_text = candidate_title.lower()
    jd_text = clean_text(post_text).lower()

    best_rule = None
    best_score = -1

    for rule in ROLE_RULES:
        score = 0
        for trigger in rule["triggers"]:
            if trigger in candidate_text:
                score += 8
            if trigger in jd_text:
                score += 2

        if score > best_score:
            best_score = score
            best_rule = rule

    if best_rule and best_score > 0:
        return best_rule

    return {
        "title": candidate_title.upper(),
        "category": "digital_marketing",
        "triggers": [],
    }


def extract_jd_skills(post_text: Any, category: str, limit: int = 30) -> List[str]:
    text = clean_text(post_text).lower()
    exact_matches = []

    for keyword, skill in JD_KEYWORD_SKILLS.items():
        if keyword in text:
            exact_matches.append(skill)

    baseline = BASE_SKILLS.get(category, BASE_SKILLS["digital_marketing"])
    skills = ordered_unique(exact_matches + baseline)

    # Keep soft skills at the end and avoid replacing technical ATS terms.
    technical = [skill for skill in skills if skill not in SOFT_SKILLS]
    return ordered_unique(technical[:26] + SOFT_SKILLS)[:limit]


# -----------------------------------------------------------------------------
# Summary and competency generation
# -----------------------------------------------------------------------------

def build_summary(role_rule: Dict[str, Any], post_text: Any, candidate: Dict[str, Any]) -> str:
    """
    Build five natural ATS-friendly sentences in ONE paragraph.
    No newline characters are returned, so Word continues immediately after each period.
    """
    category = role_rule.get("category", "digital_marketing")
    candidate_title = get_candidate_value(candidate, "title", default="Digital Marketing Specialist")
    experience = get_candidate_value(candidate, "total_experience", "experience", default="")
    experience_phrase = f"with {experience} of experience" if experience else "with hands-on experience"

    skills = extract_jd_skills(post_text, category, limit=30)
    primary = skills[:8]
    secondary = skills[8:16]

    sentence_1 = (
        f"{candidate_title} {experience_phrase} supporting SEO, paid campaigns, email marketing, "
        "content strategy, lead generation, and performance-focused digital growth."
    )

    if category == "seo":
        sentence_2 = (
            "Experienced in improving organic visibility through keyword research, on-page optimization, "
            "technical SEO, content optimization, competitor analysis, and Google Search Console."
        )
    elif category == "ppc":
        sentence_2 = (
            "Experienced in planning, monitoring, and optimizing Google Ads and PPC campaigns through keyword targeting, "
            "conversion tracking, landing-page alignment, budget control, and performance analysis."
        )
    elif category == "email_marketing":
        sentence_2 = (
            "Experienced in building and supporting email campaigns, lead-nurturing workflows, audience segmentation, "
            "marketing automation, campaign testing, and performance reporting through platforms such as Mailchimp and HubSpot."
        )
    elif category == "social_media":
        sentence_2 = (
            "Experienced in supporting organic and paid social campaigns through audience targeting, content planning, "
            "community engagement, campaign monitoring, and data-backed performance improvement."
        )
    else:
        sentence_2 = (
            "Experienced in improving online visibility, traffic quality, audience engagement, and conversion performance "
            "through structured, data-driven digital marketing execution."
        )

    sentence_3 = (
        f"Skilled in {', '.join(primary)} and related campaign execution tools used to support measurable marketing outcomes."
    )

    sentence_4 = (
        f"Brings practical exposure to {', '.join(secondary[:6])}, with a strong focus on campaign quality, "
        "reporting accuracy, and continuous optimization."
    )

    sentence_5 = (
        "Known for clear communication, cross-functional collaboration, organized execution, and aligning marketing activities "
        "with business goals without overstating experience or qualifications."
    )

    # clean_text guarantees one paragraph and one space after every period.
    return clean_text(" ".join([sentence_1, sentence_2, sentence_3, sentence_4, sentence_5]))


def build_skills(role_rule: Dict[str, Any], post_text: Any) -> str:
    category = role_rule.get("category", "digital_marketing")
    return " • ".join(extract_jd_skills(post_text, category, limit=30))


# -----------------------------------------------------------------------------
# DOCX placeholder replacement while preserving style
# -----------------------------------------------------------------------------

def replace_paragraph_text_keep_style(paragraph, new_text: Any) -> None:
    # Critical fix: remove all manual/newline characters before writing to Word.
    new_text = clean_text(new_text)

    if paragraph.runs:
        first_run = paragraph.runs[0]
        font_name = first_run.font.name
        font_size = first_run.font.size
        bold = first_run.bold
        italic = first_run.italic
        underline = first_run.underline

        first_run.text = new_text
        first_run.font.name = font_name
        first_run.font.size = font_size
        first_run.bold = bold
        first_run.italic = italic
        first_run.underline = underline

        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def replace_in_paragraphs(paragraphs, replacements: Dict[str, str]) -> None:
    for paragraph in paragraphs:
        original_text = paragraph.text
        if not original_text:
            continue

        updated_text = original_text
        changed = False
        for placeholder, replacement in replacements.items():
            if placeholder in updated_text:
                updated_text = updated_text.replace(placeholder, str(replacement))
                changed = True

        if changed:
            replace_paragraph_text_keep_style(paragraph, updated_text)


def replace_in_tables(tables, replacements: Dict[str, str]) -> None:
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                replace_in_paragraphs(cell.paragraphs, replacements)
                replace_in_tables(cell.tables, replacements)


def replace_text_in_docx(doc: Document, replacements: Dict[str, str]) -> None:
    replace_in_paragraphs(doc.paragraphs, replacements)
    replace_in_tables(doc.tables, replacements)

    for section in doc.sections:
        replace_in_paragraphs(section.header.paragraphs, replacements)
        replace_in_tables(section.header.tables, replacements)
        replace_in_paragraphs(section.footer.paragraphs, replacements)
        replace_in_tables(section.footer.tables, replacements)


def collect_unreplaced_placeholders(doc: Document) -> List[str]:
    texts = []

    def collect_paragraphs(paragraphs) -> None:
        texts.extend(paragraph.text for paragraph in paragraphs)

    collect_paragraphs(doc.paragraphs)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                collect_paragraphs(cell.paragraphs)

    for section in doc.sections:
        collect_paragraphs(section.header.paragraphs)
        collect_paragraphs(section.footer.paragraphs)

    return re.findall(r"\{\{[^}]+\}\}", "\n".join(texts))


# -----------------------------------------------------------------------------
# Public function used by existing main.py
# -----------------------------------------------------------------------------

def customize_resume(post_text: Any, lead_index: int) -> str:
    OUTPUT_RESUME_DIR.mkdir(parents=True, exist_ok=True)

    candidate = load_candidate_data()
    name = get_candidate_value(candidate, "name", "full_name", default="Candidate")
    safe_name = safe_filename(name)

    template_path = detect_resume_template(candidate)
    custom_docx = OUTPUT_RESUME_DIR / f"{safe_name}_Resume_{lead_index:03d}.docx"
    custom_pdf = OUTPUT_RESUME_DIR / f"{safe_name}_Resume_{lead_index:03d}.pdf"

    role_rule = detect_role(post_text, candidate)
    title = get_candidate_value(candidate, "title", default=role_rule.get("title", "Digital Marketing Specialist"))
    summary = build_summary(role_rule, post_text, candidate)
    skills = build_skills(role_rule, post_text)

    replacements = {
        "{{NAME}}": name,
        "{{TITLE}}": title,
        "{{EMAIL}}": get_candidate_value(candidate, "email", default=""),
        "{{PHONE}}": get_candidate_value(candidate, "phone", "mobile", default=""),
        "{{LOCATION}}": get_candidate_value(candidate, "location", "current_location", default=""),
        "{{LINKEDIN}}": get_candidate_value(candidate, "linkedin", "linkedin_url", default=""),
        "{{SUMMARY}}": summary,
        "{{SKILLS}}": skills,
        "{{CORE_COMPETENCIES}}": skills,
    }

    doc = Document(template_path)
    replace_text_in_docx(doc, replacements)

    leftovers = collect_unreplaced_placeholders(doc)
    if leftovers:
        print(f"⚠️ Warning: Unreplaced placeholders found: {leftovers}")

    doc.save(custom_docx)

    try:
        convert(str(custom_docx), str(custom_pdf))
    except Exception as error:
        raise RuntimeError(
            f"DOCX was created at {custom_docx}, but PDF conversion failed: {error}"
        ) from error

    if not custom_pdf.exists():
        raise FileNotFoundError(f"PDF was not created: {custom_pdf}")

    return str(custom_pdf)


if __name__ == "__main__":
    sample_jd = (
        "Hiring SEO Specialist with keyword research, on-page SEO, technical SEO, Google Search Console, "
        "Google Analytics, GA4, competitor analysis, content optimization, landing pages and campaign reporting."
    )
    print(customize_resume(sample_jd, 1))
